"""
Export service for generating meeting reports in various formats.
"""

import json
import tempfile
from typing import Dict, List, Optional, Any
from pathlib import Path
from datetime import datetime
import io

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors

from app.core.logging import get_logger
from app.models.meeting import Meeting

logger = get_logger(__name__)


class ExportService:
    """
    Service for exporting meeting data to various formats.
    """
    
    def __init__(self):
        self.logger = logger
    
    def export_to_json(
        self,
        meeting: Meeting,
        include_transcript: bool = True,
        include_summary: bool = True,
        include_action_items: bool = True,
        include_key_decisions: bool = True,
        include_speaker_segments: bool = False
    ) -> str:
        """Export meeting data to JSON format."""
        try:
            export_data = {
                "meeting_info": {
                    "id": meeting.id,
                    "title": meeting.title,
                    "description": meeting.description,
                    "created_at": meeting.created_at.isoformat(),
                    "duration": meeting.duration,
                    "language_detected": meeting.language_detected,
                    "confidence_score": meeting.confidence_score,
                    "tags": meeting.tags,
                    "context": meeting.context
                },
                "processing_info": {
                    "asr_provider": meeting.asr_provider.value if meeting.asr_provider else None,
                    "asr_model": meeting.asr_model,
                    "llm_provider": meeting.llm_provider.value if meeting.llm_provider else None,
                    "llm_model": meeting.llm_model,
                    "status": meeting.status.value if meeting.status else None,
                    "num_speakers": meeting.num_speakers
                }
            }
            
            if include_transcript and meeting.transcript:
                export_data["transcript"] = meeting.transcript
            
            if include_summary and meeting.summaries:
                export_data["summaries"] = meeting.summaries
            
            if include_action_items and meeting.action_items:
                export_data["action_items"] = meeting.action_items
            
            if include_key_decisions and meeting.key_decisions:
                export_data["key_decisions"] = meeting.key_decisions
            
            if meeting.risks:
                export_data["risks"] = meeting.risks
            
            if meeting.open_questions:
                export_data["open_questions"] = meeting.open_questions
            
            if meeting.keywords:
                export_data["keywords"] = meeting.keywords
            
            if include_speaker_segments and meeting.speaker_segments:
                export_data["speaker_segments"] = meeting.speaker_segments
            
            return json.dumps(export_data, indent=2, ensure_ascii=False)
            
        except Exception as e:
            self.logger.error(f"Error exporting to JSON: {e}")
            raise
    
    def export_to_markdown(
        self,
        meeting: Meeting,
        include_transcript: bool = True,
        include_summary: bool = True,
        include_action_items: bool = True,
        include_key_decisions: bool = True,
        include_speaker_segments: bool = False
    ) -> str:
        """Export meeting data to Markdown format."""
        try:
            lines = []
            
            # Header
            lines.append(f"# {meeting.title}")
            lines.append("")
            
            if meeting.description:
                lines.append(f"**Description:** {meeting.description}")
                lines.append("")
            
            # Meeting Info
            lines.append("## Meeting Information")
            lines.append("")
            lines.append(f"- **Date:** {meeting.created_at.strftime('%Y-%m-%d %H:%M:%S')}")
            if meeting.duration:
                lines.append(f"- **Duration:** {meeting.duration:.1f} seconds")
            if meeting.language_detected:
                lines.append(f"- **Language:** {meeting.language_detected}")
            if meeting.confidence_score:
                lines.append(f"- **Confidence Score:** {meeting.confidence_score:.2f}")
            if meeting.num_speakers:
                lines.append(f"- **Number of Speakers:** {meeting.num_speakers}")
            if meeting.tags:
                lines.append(f"- **Tags:** {', '.join(meeting.tags)}")
            lines.append("")
            
            # Context
            if meeting.context:
                lines.append("## Context")
                lines.append("")
                lines.append(meeting.context)
                lines.append("")
            
            # Executive Summary
            if include_summary and meeting.summaries:
                if meeting.summaries.get('executive_summary'):
                    lines.append("## Executive Summary")
                    lines.append("")
                    lines.append(meeting.summaries['executive_summary'])
                    lines.append("")
            
            # Key Decisions
            if include_key_decisions and meeting.key_decisions:
                lines.append("## Key Decisions")
                lines.append("")
                for i, decision in enumerate(meeting.key_decisions, 1):
                    lines.append(f"### {i}. {decision.get('description', 'N/A')}")
                    if decision.get('decision_maker'):
                        lines.append(f"**Decision Maker:** {decision['decision_maker']}")
                    if decision.get('rationale'):
                        lines.append(f"**Rationale:** {decision['rationale']}")
                    lines.append("")
            
            # Action Items
            if include_action_items and meeting.action_items:
                lines.append("## Action Items")
                lines.append("")
                for i, item in enumerate(meeting.action_items, 1):
                    lines.append(f"### {i}. {item.get('description', 'N/A')}")
                    if item.get('owner'):
                        lines.append(f"**Owner:** {item['owner']}")
                    if item.get('deadline'):
                        lines.append(f"**Deadline:** {item['deadline']}")
                    if item.get('priority'):
                        lines.append(f"**Priority:** {item['priority']}")
                    lines.append("")
            
            # Risks
            if meeting.risks:
                lines.append("## Risks and Concerns")
                lines.append("")
                for i, risk in enumerate(meeting.risks, 1):
                    lines.append(f"### {i}. {risk.get('description', 'N/A')}")
                    if risk.get('impact'):
                        lines.append(f"**Impact:** {risk['impact']}")
                    if risk.get('probability'):
                        lines.append(f"**Probability:** {risk['probability']}")
                    if risk.get('mitigation'):
                        lines.append(f"**Mitigation:** {risk['mitigation']}")
                    lines.append("")
            
            # Open Questions
            if meeting.open_questions:
                lines.append("## Open Questions")
                lines.append("")
                for i, question in enumerate(meeting.open_questions, 1):
                    lines.append(f"{i}. {question.get('question', 'N/A')}")
                    if question.get('assigned_to'):
                        lines.append(f"   - **Assigned to:** {question['assigned_to']}")
                lines.append("")
            
            # Keywords
            if meeting.keywords:
                lines.append("## Keywords")
                lines.append("")
                lines.append(", ".join(meeting.keywords))
                lines.append("")
            
            # Detailed Summary
            if include_summary and meeting.summaries and meeting.summaries.get('detailed_summary'):
                lines.append("## Detailed Summary")
                lines.append("")
                lines.append(meeting.summaries['detailed_summary'])
                lines.append("")
            
            # Transcript
            if include_transcript and meeting.transcript:
                lines.append("## Transcript")
                lines.append("")
                lines.append("```")
                lines.append(meeting.transcript)
                lines.append("```")
                lines.append("")
            
            return "\n".join(lines)
            
        except Exception as e:
            self.logger.error(f"Error exporting to Markdown: {e}")
            raise
    
    def export_to_docx(
        self,
        meeting: Meeting,
        include_transcript: bool = True,
        include_summary: bool = True,
        include_action_items: bool = True,
        include_key_decisions: bool = True,
        include_speaker_segments: bool = False
    ) -> bytes:
        """Export meeting data to DOCX format."""
        try:
            doc = Document()
            
            # Title
            title = doc.add_heading(meeting.title, 0)
            
            # Meeting Information
            doc.add_heading('Meeting Information', level=1)
            
            info_table = doc.add_table(rows=0, cols=2)
            info_table.style = 'Table Grid'
            
            # Add meeting info rows
            info_items = [
                ('Date', meeting.created_at.strftime('%Y-%m-%d %H:%M:%S')),
                ('Duration', f"{meeting.duration:.1f} seconds" if meeting.duration else "N/A"),
                ('Language', meeting.language_detected or "N/A"),
                ('Confidence Score', f"{meeting.confidence_score:.2f}" if meeting.confidence_score else "N/A"),
                ('Number of Speakers', str(meeting.num_speakers) if meeting.num_speakers else "N/A"),
                ('Tags', ', '.join(meeting.tags) if meeting.tags else "N/A")
            ]
            
            for label, value in info_items:
                row_cells = info_table.add_row().cells
                row_cells[0].text = label
                row_cells[1].text = value
            
            # Context
            if meeting.context:
                doc.add_heading('Context', level=1)
                doc.add_paragraph(meeting.context)
            
            # Executive Summary
            if include_summary and meeting.summaries and meeting.summaries.get('executive_summary'):
                doc.add_heading('Executive Summary', level=1)
                doc.add_paragraph(meeting.summaries['executive_summary'])
            
            # Key Decisions
            if include_key_decisions and meeting.key_decisions:
                doc.add_heading('Key Decisions', level=1)
                for i, decision in enumerate(meeting.key_decisions, 1):
                    doc.add_heading(f"{i}. {decision.get('description', 'N/A')}", level=2)
                    if decision.get('decision_maker'):
                        doc.add_paragraph(f"Decision Maker: {decision['decision_maker']}")
                    if decision.get('rationale'):
                        doc.add_paragraph(f"Rationale: {decision['rationale']}")
            
            # Action Items
            if include_action_items and meeting.action_items:
                doc.add_heading('Action Items', level=1)
                
                action_table = doc.add_table(rows=1, cols=4)
                action_table.style = 'Table Grid'
                
                # Header row
                header_cells = action_table.rows[0].cells
                header_cells[0].text = 'Description'
                header_cells[1].text = 'Owner'
                header_cells[2].text = 'Deadline'
                header_cells[3].text = 'Priority'
                
                for item in meeting.action_items:
                    row_cells = action_table.add_row().cells
                    row_cells[0].text = item.get('description', 'N/A')
                    row_cells[1].text = item.get('owner', 'N/A')
                    row_cells[2].text = item.get('deadline', 'N/A')
                    row_cells[3].text = item.get('priority', 'N/A')
            
            # Risks
            if meeting.risks:
                doc.add_heading('Risks and Concerns', level=1)
                for i, risk in enumerate(meeting.risks, 1):
                    doc.add_heading(f"{i}. {risk.get('description', 'N/A')}", level=2)
                    doc.add_paragraph(f"Impact: {risk.get('impact', 'N/A')}")
                    doc.add_paragraph(f"Probability: {risk.get('probability', 'N/A')}")
                    if risk.get('mitigation'):
                        doc.add_paragraph(f"Mitigation: {risk['mitigation']}")
            
            # Open Questions
            if meeting.open_questions:
                doc.add_heading('Open Questions', level=1)
                for i, question in enumerate(meeting.open_questions, 1):
                    para = doc.add_paragraph(f"{i}. {question.get('question', 'N/A')}")
                    if question.get('assigned_to'):
                        doc.add_paragraph(f"   Assigned to: {question['assigned_to']}")
            
            # Keywords
            if meeting.keywords:
                doc.add_heading('Keywords', level=1)
                doc.add_paragraph(', '.join(meeting.keywords))
            
            # Detailed Summary
            if include_summary and meeting.summaries and meeting.summaries.get('detailed_summary'):
                doc.add_heading('Detailed Summary', level=1)
                doc.add_paragraph(meeting.summaries['detailed_summary'])
            
            # Transcript
            if include_transcript and meeting.transcript:
                doc.add_heading('Transcript', level=1)
                doc.add_paragraph(meeting.transcript)
            
            # Save to bytes
            doc_buffer = io.BytesIO()
            doc.save(doc_buffer)
            doc_buffer.seek(0)
            
            return doc_buffer.getvalue()
            
        except Exception as e:
            self.logger.error(f"Error exporting to DOCX: {e}")
            raise
    
    def export_to_pdf(
        self,
        meeting: Meeting,
        include_transcript: bool = True,
        include_summary: bool = True,
        include_action_items: bool = True,
        include_key_decisions: bool = True,
        include_speaker_segments: bool = False
    ) -> bytes:
        """Export meeting data to PDF format."""
        try:
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            
            # Styles
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                spaceAfter=30,
            )
            
            # Build content
            story = []
            
            # Title
            story.append(Paragraph(meeting.title, title_style))
            story.append(Spacer(1, 12))
            
            # Meeting Information
            story.append(Paragraph("Meeting Information", styles['Heading1']))
            
            info_data = [
                ['Date', meeting.created_at.strftime('%Y-%m-%d %H:%M:%S')],
                ['Duration', f"{meeting.duration:.1f} seconds" if meeting.duration else "N/A"],
                ['Language', meeting.language_detected or "N/A"],
                ['Confidence Score', f"{meeting.confidence_score:.2f}" if meeting.confidence_score else "N/A"],
                ['Number of Speakers', str(meeting.num_speakers) if meeting.num_speakers else "N/A"],
                ['Tags', ', '.join(meeting.tags) if meeting.tags else "N/A"]
            ]
            
            info_table = Table(info_data, colWidths=[2*inch, 4*inch])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 14),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(info_table)
            story.append(Spacer(1, 12))
            
            # Context
            if meeting.context:
                story.append(Paragraph("Context", styles['Heading1']))
                story.append(Paragraph(meeting.context, styles['Normal']))
                story.append(Spacer(1, 12))
            
            # Executive Summary
            if include_summary and meeting.summaries and meeting.summaries.get('executive_summary'):
                story.append(Paragraph("Executive Summary", styles['Heading1']))
                story.append(Paragraph(meeting.summaries['executive_summary'], styles['Normal']))
                story.append(Spacer(1, 12))
            
            # Key Decisions
            if include_key_decisions and meeting.key_decisions:
                story.append(Paragraph("Key Decisions", styles['Heading1']))
                for i, decision in enumerate(meeting.key_decisions, 1):
                    story.append(Paragraph(f"{i}. {decision.get('description', 'N/A')}", styles['Heading2']))
                    if decision.get('decision_maker'):
                        story.append(Paragraph(f"<b>Decision Maker:</b> {decision['decision_maker']}", styles['Normal']))
                    if decision.get('rationale'):
                        story.append(Paragraph(f"<b>Rationale:</b> {decision['rationale']}", styles['Normal']))
                    story.append(Spacer(1, 6))
            
            # Action Items
            if include_action_items and meeting.action_items:
                story.append(Paragraph("Action Items", styles['Heading1']))
                
                action_data = [['Description', 'Owner', 'Deadline', 'Priority']]
                for item in meeting.action_items:
                    action_data.append([
                        item.get('description', 'N/A'),
                        item.get('owner', 'N/A'),
                        item.get('deadline', 'N/A'),
                        item.get('priority', 'N/A')
                    ])
                
                action_table = Table(action_data)
                action_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                
                story.append(action_table)
                story.append(Spacer(1, 12))
            
            # Add other sections similarly...
            
            # Transcript
            if include_transcript and meeting.transcript:
                story.append(Paragraph("Transcript", styles['Heading1']))
                # Split long transcript into smaller paragraphs
                transcript_lines = meeting.transcript.split('\n')
                for line in transcript_lines[:50]:  # Limit for PDF
                    if line.strip():
                        story.append(Paragraph(line, styles['Normal']))
            
            # Build PDF
            doc.build(story)
            
            buffer.seek(0)
            return buffer.getvalue()
            
        except Exception as e:
            self.logger.error(f"Error exporting to PDF: {e}")
            raise
    
    def export_meeting(
        self,
        meeting: Meeting,
        format_type: str,
        include_transcript: bool = True,
        include_summary: bool = True,
        include_action_items: bool = True,
        include_key_decisions: bool = True,
        include_speaker_segments: bool = False
    ) -> tuple[bytes, str]:
        """
        Export meeting in specified format.
        
        Returns:
            (content_bytes, content_type)
        """
        try:
            if format_type.lower() == 'json':
                content = self.export_to_json(
                    meeting, include_transcript, include_summary,
                    include_action_items, include_key_decisions, include_speaker_segments
                )
                return content.encode('utf-8'), 'application/json'
            
            elif format_type.lower() == 'markdown':
                content = self.export_to_markdown(
                    meeting, include_transcript, include_summary,
                    include_action_items, include_key_decisions, include_speaker_segments
                )
                return content.encode('utf-8'), 'text/markdown'
            
            elif format_type.lower() == 'docx':
                content = self.export_to_docx(
                    meeting, include_transcript, include_summary,
                    include_action_items, include_key_decisions, include_speaker_segments
                )
                return content, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            
            elif format_type.lower() == 'pdf':
                content = self.export_to_pdf(
                    meeting, include_transcript, include_summary,
                    include_action_items, include_key_decisions, include_speaker_segments
                )
                return content, 'application/pdf'
            
            else:
                raise ValueError(f"Unsupported export format: {format_type}")
                
        except Exception as e:
            self.logger.error(f"Error exporting meeting in format {format_type}: {e}")
            raise


# Create global service instance
export_service = ExportService()